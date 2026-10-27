"""
Text extraction, chunking and passage selection for user-uploaded files.

This is separate from the permanent knowledge base: uploads are only ever
used to answer their owner's questions in the same chat session.

Supported: PDF (text layer; scanned pages through provider OCR when
enabled), DOCX (paragraphs and tables), TXT/Markdown, CSV and XLSX (rows
keep their column headers). Nothing in a document is executed: macros are
rejected at detection, formulas are read as their cached values, and the
text is always passed to the model as untrusted data.
"""

import csv
import io
import math
import re
import time
from collections import Counter
from dataclasses import dataclass, field

from apps.aiChat import ai_config
from apps.aiChat.attachments.errors import AttachmentError
from apps.aiChat.rag.document_loader import split_document

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TXT = "text/plain"
CSV = "text/csv"

# Below this many characters per page a PDF is treated as scanned.
_SCANNED_CHARS_PER_PAGE = 25


@dataclass
class Section:
    text: str
    location: str  # e.g. "page 3", "sheet Schedule", "part 2"


@dataclass
class ExtractedDocument:
    name: str
    mime_type: str
    sections: list[Section] = field(default_factory=list)
    page_count: int | None = None
    pages_read: int | None = None
    truncated: bool = False
    scanned: bool = False
    ocr_used: bool = False

    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.sections)

    def to_dict(self) -> dict:
        return {
            "name": self.name,
            "mime_type": self.mime_type,
            "page_count": self.page_count,
            "pages_read": self.pages_read,
            "truncated": self.truncated,
            "scanned": self.scanned,
            "ocr_used": self.ocr_used,
            "sections": [
                {"text": s.text, "location": s.location} for s in self.sections
            ],
        }

    @classmethod
    def from_dict(cls, value: dict) -> "ExtractedDocument":
        return cls(
            name=value.get("name", ""),
            mime_type=value.get("mime_type", ""),
            sections=[
                Section(s.get("text", ""), s.get("location", ""))
                for s in value.get("sections", [])
                if isinstance(s, dict)
            ],
            page_count=value.get("page_count"),
            pages_read=value.get("pages_read"),
            truncated=bool(value.get("truncated")),
            scanned=bool(value.get("scanned")),
            ocr_used=bool(value.get("ocr_used")),
        )


class _Budget:
    """Caps extracted characters and extraction time."""

    def __init__(self):
        self.chars_left = ai_config.max_extracted_chars()
        self.deadline = time.monotonic() + ai_config.document_extract_seconds()
        self.exhausted = False

    def take(self, text: str) -> str:
        if self.exhausted:
            return ""
        if time.monotonic() > self.deadline or self.chars_left <= 0:
            self.exhausted = True
            return ""
        if len(text) > self.chars_left:
            text = text[: self.chars_left]
            self.exhausted = True
        self.chars_left -= len(text)
        return text


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


# ---- Extractors -------------------------------------------------------------

def _extract_pdf(data: bytes, doc: ExtractedDocument) -> None:
    from pypdf import PdfReader
    from pypdf.errors import FileNotDecryptedError, PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # Some PDFs are "encrypted" with an empty user password only.
            try:
                if not reader.decrypt(""):
                    raise AttachmentError("encrypted_document")
            except (FileNotDecryptedError, NotImplementedError) as exc:
                raise AttachmentError("encrypted_document") from exc
        pages = reader.pages
        doc.page_count = len(pages)
    except AttachmentError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as exc:
        raise AttachmentError("corrupt_document") from exc

    limit = min(doc.page_count, ai_config.max_document_pages())
    doc.truncated = doc.page_count > limit
    budget = _Budget()
    raw_chars = 0
    read = 0
    for index in range(limit):
        try:
            text = pages[index].extract_text() or ""
        except Exception:  # noqa: BLE001 - one bad page must not fail all
            text = ""
        read += 1
        text = _clean(text)
        raw_chars += len(text)
        kept = budget.take(text)
        if kept:
            doc.sections.append(Section(kept, f"page {index + 1}"))
        if budget.exhausted:
            doc.truncated = True
            break
    doc.pages_read = read
    doc.scanned = read > 0 and raw_chars < _SCANNED_CHARS_PER_PAGE * read


def _extract_docx(data: bytes, doc: ExtractedDocument) -> None:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises many types
        raise AttachmentError("corrupt_document") from exc

    budget = _Budget()
    part = 1
    lines: list[str] = []

    def flush():
        nonlocal part, lines
        text = budget.take(_clean("\n".join(lines)))
        if text:
            doc.sections.append(Section(text, f"part {part}"))
            part += 1
        lines = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading") and lines:
            flush()
        lines.append(text)
        if budget.exhausted:
            break
    for table_number, table in enumerate(document.tables, start=1):
        if budget.exhausted:
            break
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            lines.append(f"Table {table_number}:\n" + "\n".join(rows))
    flush()
    doc.truncated = budget.exhausted


def _decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AttachmentError("corrupt_document") from exc


def _extract_txt(data: bytes, doc: ExtractedDocument) -> None:
    budget = _Budget()
    text = budget.take(_clean(_decode_text(data)))
    if text:
        doc.sections.append(Section(text, "text"))
    doc.truncated = budget.exhausted


def _rows_as_text(rows, location: str, budget: _Budget) -> list[Section]:
    """Rows become 'Row N: Header=value; ...' so values keep their columns."""
    max_rows = ai_config.max_sheet_rows()
    sections: list[Section] = []
    header: list[str] | None = None
    lines: list[str] = []
    for number, row in enumerate(rows, start=1):
        cells = ["" if v is None else str(v).strip() for v in row][:50]
        if not any(cells):
            continue
        if header is None:
            header = [c or f"Column {i + 1}" for i, c in enumerate(cells)]
            lines.append("Columns: " + " | ".join(header))
            continue
        if number > max_rows:
            budget.exhausted = True
            break
        pairs = [
            f"{header[i] if i < len(header) else f'Column {i + 1}'}={value}"
            for i, value in enumerate(cells)
            if value
        ]
        line = budget.take(f"Row {number}: " + "; ".join(pairs))
        if not line:
            break
        lines.append(line)
    if len(lines) > 1 or (lines and header):
        sections.append(Section("\n".join(lines), location))
    return sections


def _extract_csv(data: bytes, doc: ExtractedDocument) -> None:
    text = _decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    budget = _Budget()
    try:
        doc.sections.extend(
            _rows_as_text(csv.reader(io.StringIO(text), dialect), "table", budget)
        )
    except csv.Error as exc:
        raise AttachmentError("corrupt_document") from exc
    doc.truncated = budget.exhausted


def _extract_xlsx(data: bytes, doc: ExtractedDocument) -> None:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True,
        )
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types
        raise AttachmentError("corrupt_document") from exc
    budget = _Budget()
    try:
        for sheet in workbook.worksheets:
            if budget.exhausted:
                break
            doc.sections.extend(
                _rows_as_text(
                    sheet.iter_rows(values_only=True),
                    f"sheet {sheet.title}",
                    budget,
                )
            )
    finally:
        workbook.close()
    doc.truncated = budget.exhausted


_EXTRACTORS = {
    PDF: _extract_pdf,
    DOCX: _extract_docx,
    XLSX: _extract_xlsx,
    TXT: _extract_txt,
    CSV: _extract_csv,
}


def extract_document(data: bytes, mime_type: str, name: str) -> ExtractedDocument:
    """Extract text or raise `AttachmentError`. Never runs embedded code."""
    if len(data) > ai_config.max_document_bytes():
        raise AttachmentError(
            "file_too_large",
            limit=ai_config.max_document_bytes() // (1024 * 1024),
        )
    extractor = _EXTRACTORS.get(mime_type)
    if extractor is None:
        raise AttachmentError("unsupported_file")
    doc = ExtractedDocument(name=name, mime_type=mime_type)
    extractor(data, doc)
    return doc


def apply_ocr_pages(doc: ExtractedDocument, pages: list[dict]) -> None:
    """Replace a scanned PDF's sections with provider OCR text."""
    budget = _Budget()
    sections = []
    for item in pages:
        if not isinstance(item, dict):
            continue
        text = budget.take(_clean(str(item.get("text") or "")))
        if text:
            sections.append(Section(text, f"page {item.get('page', '?')}"))
        if budget.exhausted:
            doc.truncated = True
            break
    doc.sections = sections
    doc.ocr_used = True


# ---- Chunking and selection -------------------------------------------------

@dataclass(frozen=True)
class Chunk:
    index: int
    text: str
    location: str


def chunk_document(doc: ExtractedDocument) -> list[Chunk]:
    chunks: list[Chunk] = []
    for section in doc.sections:
        for piece in split_document(section.text, chunk_size=1200, chunk_overlap=150):
            chunks.append(Chunk(len(chunks), piece, section.location))
    return chunks


_SUMMARY_WORDS = (
    "summar", "overview", "gist", "explain this", "explain the", "what is this",
    "what does this", "tl;dr", "सारांश", "सार", "थोडक्यात", "संक्षेप",
    "समझाइए", "समझाओ", "स्पष्ट करा",
)
_WORD = re.compile(r"[0-9a-zऀ-ॿ]+")
_STOP = {
    "the", "and", "for", "with", "this", "that", "what", "does", "about", "from",
    "are", "was", "were", "which", "when", "where", "how", "document", "file",
    "pdf", "say", "says", "tell", "please", "notice", "है", "का", "की", "के",
    "में", "क्या", "आहे", "काय", "या", "हा", "ही",
}


def _terms(text: str) -> list[str]:
    return [
        word for word in _WORD.findall(text.lower())
        if word not in _STOP and (len(word) >= 3 or not word.isascii())
    ]


def _spread(chunks: list[Chunk], count: int) -> list[Chunk]:
    if len(chunks) <= count:
        return list(chunks)
    step = (len(chunks) - 1) / max(1, count - 1)
    picked = sorted({round(i * step) for i in range(count)})
    return [chunks[i] for i in picked]


def select_chunks(chunks: list[Chunk], question: str) -> list[Chunk]:
    """
    Pick the passages to send to the model within the context budget.

    Summaries get passages spread across the whole document; specific
    questions get the passages sharing the most (rarer) terms with the
    question. When nothing matches, fall back to the spread so the model
    can still say whether the document covers the question.
    """
    if not chunks:
        return []
    budget = ai_config.document_context_chars()
    max_count = max(1, ai_config.max_document_chunks())
    lowered = (question or "").lower()

    if sum(len(c.text) for c in chunks) <= budget:
        return list(chunks)
    if not lowered.strip() or any(w in lowered for w in _SUMMARY_WORDS):
        ranked = _spread(chunks, max_count)
    else:
        query = set(_terms(question))
        frequency = Counter()
        chunk_terms = []
        for chunk in chunks:
            terms = Counter(_terms(chunk.text))
            chunk_terms.append(terms)
            frequency.update(terms.keys())
        total = len(chunks)
        scored = []
        for chunk, terms in zip(chunks, chunk_terms):
            score = sum(
                (1 + math.log(terms[t])) * math.log(1 + total / frequency[t])
                for t in query
                if terms.get(t)
            )
            scored.append((score, chunk))
        best = [c for s, c in sorted(scored, key=lambda x: -x[0]) if s > 0]
        ranked = best[:max_count] if best else _spread(chunks, max_count)

    selected, used = [], 0
    for chunk in ranked:
        if len(selected) >= max_count or used + len(chunk.text) > budget:
            if selected:
                break
        selected.append(chunk)
        used += len(chunk.text)
    return sorted(selected, key=lambda c: c.index)
