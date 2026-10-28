"""Shared fixtures for Guru AI tests. No test makes a live provider call."""

import io
import os
import shutil
import tempfile
import wave
from array import array
from datetime import date, timedelta
from decimal import Decimal
from unittest import mock

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from PIL import Image
from rest_framework.test import APIClient

from apps.aiChat.models import ChatMessage, ChatSession
from apps.authentication.models import User
from apps.crowd.models import Incident, Zone, ZoneCrowdSnapshot
from apps.events.models.EventModel import Event
from apps.temples.models import Temple

SEND_URL = "/api/chat/chat-send/"
HISTORY_URL = "/api/chat/chat-history/"


# ---- files ----------------------------------------------------------------------

def png_bytes(size=(64, 48), color=(255, 120, 0)) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


def jpeg_with_exif() -> bytes:
    image = Image.new("RGB", (40, 30), (10, 120, 200))
    exif = Image.Exif()
    exif[0x010F] = "PhoneMaker"  # Make
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", exif=exif.tobytes())
    return buffer.getvalue()


def wav_bytes(seconds=1.0, amplitude=8000, rate=16000) -> bytes:
    frames = int(seconds * rate)
    samples = array("h", [
        amplitude if (i // 20) % 2 else -amplitude for i in range(frames)
    ])
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(rate)
        wav.writeframes(samples.tobytes())
    return buffer.getvalue()


def _pdf_string(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_bytes(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    objects: list[bytes] = []
    page_ids = []
    font_id = 3
    next_id = 4
    page_objects = []
    for text in pages:
        content = (
            f"BT /F1 12 Tf 72 720 Td ({_pdf_string(text)}) Tj ET".encode()
            if text else b""
        )
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        page_ids.append(page_id)
        page_objects.append((page_id, (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {content_id} 0 R >>"
        ).encode()))
        page_objects.append((
            content_id,
            b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        ))
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    numbered = [
        (1, b"<< /Type /Catalog /Pages 2 0 R >>"),
        (2, f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()),
        (3, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"),
        *page_objects,
    ]
    numbered.sort()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for number, body in numbered:
        offsets[number] = out.tell()
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    total = len(numbered) + 1
    out.write(f"xref\n0 {total}\n0000000000 65535 f \n".encode())
    for number in range(1, total):
        out.write(f"{offsets[number]:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {total} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    del objects
    return out.getvalue()


def encrypted_pdf_bytes() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="secret", owner_password="owner")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(rows: list[list]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Schedule"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(name: str, data: bytes, content_type: str) -> SimpleUploadedFile:
    return SimpleUploadedFile(name, data, content_type=content_type)


# ---- data -------------------------------------------------------------------------

def make_temple(name="Trimbakeshwar Jyotirling Mandir", lat="19.932400", lng="73.530800", **extra):
    return Temple.objects.create(
        name=name, latitude=Decimal(lat), longitude=Decimal(lng), city="Nashik",
        description=extra.pop("description", "Jyotirlinga temple of Lord Shiva."),
        **extra,
    )


def make_event(name="Amrit Snan", days=10, **extra):
    values = dict(
        event_name=name,
        event_category="SNAN",
        event_subtype="AMRIT_SNAN",
        event_date=date.today() + timedelta(days=days),
        location_type="NASHIK",
        location_name="Ramkund",
        data_source="GOVT",
        is_verified=True,
        status="PUBLISHED",
        show_in_app=True,
    )
    values.update(extra)
    return Event.objects.create(**values)


def make_zone(code="RAMKUND_01", name="Ramkund", lat="19.997500", lng="73.791000", **extra):
    return Zone.objects.create(
        code=code, name=name, city="Nashik", latitude=Decimal(lat),
        longitude=Decimal(lng), radius_m=300, capacity=1000, **extra,
    )


def make_snapshot(zone, minutes_ago=5, status="MEDIUM", occupancy=55.0):
    return ZoneCrowdSnapshot.objects.create(
        zone=zone, total_people=int(occupancy * 10), occupancy_percent=occupancy,
        status=status, calculated_at=timezone.now() - timedelta(minutes=minutes_ago),
    )


def make_incident(zone, title="Heavy traffic near the ghat", **extra):
    values = dict(
        zone=zone, incident_type="TRAFFIC", severity="HIGH", title=title,
        description="Vehicles diverted via the ring road.",
        start_time=timezone.now() - timedelta(minutes=30), is_active=True,
    )
    values.update(extra)
    return Incident.objects.create(**values)


# ---- base test case ----------------------------------------------------------------

def _no_live_model(*args, **kwargs):
    raise AssertionError("A test reached the real Gemini client.")


class GuruTestCase(TestCase):
    """Authenticated client, temporary media folder, no live model calls."""

    env: dict = {}

    def setUp(self):
        cache.clear()
        self.media = tempfile.mkdtemp(prefix="guru_media_")
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        media = override_settings(MEDIA_ROOT=self.media)
        media.enable()
        self.addCleanup(media.disable)

        environment = {
            "GEMINI_API_KEY": "test-key",
            "AI_MEDIA_PIPELINE": "native",
            "AI_VOICE_REPLY_ENABLED": "false",
            "AI_CHAT_RATE": "1000/min",
            **self.env,
        }
        patcher = mock.patch.dict(os.environ, environment)
        patcher.start()
        self.addCleanup(patcher.stop)

        # Fail loudly if anything reaches a real provider or the local
        # vector store; knowledge search returns nothing unless a test
        # supplies results with mock_knowledge().
        for target in (
            "apps.aiChat.services.gemini_service._chat_model",
            "apps.aiChat.rag.knowledge_service.get_embeddings",
            "apps.aiChat.rag.knowledge_service.qdrant_session",
        ):
            guard = mock.patch(target, side_effect=_no_live_model)
            guard.start()
            self.addCleanup(guard.stop)
        self.mock_knowledge()

        self.user = User.objects.create_user(phone_number="9000000001", email="a@example.com")
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    # The model as the composer sees it.
    def mock_model(self, reply="Mocked answer.", side_effect=None):
        patcher = mock.patch(
            "apps.aiChat.services.answer_composer.invoke_gemini",
            return_value=reply, side_effect=side_effect,
        )
        model = patcher.start()
        self.addCleanup(patcher.stop)
        return model

    def mock_knowledge(self, results=None):
        patcher = mock.patch(
            "apps.aiChat.graph.search_documents", return_value=results or [],
        )
        search = patcher.start()
        self.addCleanup(patcher.stop)
        return search

    def send(self, message=None, file=None, **fields):
        data = {k: v for k, v in fields.items() if v is not None}
        if message is not None:
            data["message"] = message
        if file is not None:
            data["file"] = file
            return self.client.post(SEND_URL, data, format="multipart")
        return self.client.post(SEND_URL, data, format="json")

    def reply(self, response) -> dict:
        self.assertEqual(response.status_code, 200, response.content[:500])
        return response.json()["data"]["messages"][1]

    @staticmethod
    def prompt_text(model) -> str:
        content = model.call_args.args[0]
        return "\n".join(b["text"] for b in content if b.get("type") == "text")

    @staticmethod
    def system_prompt(model) -> str:
        return model.call_args.kwargs["system"]

    def session(self) -> ChatSession:
        return ChatSession.objects.get(user=self.user)

    def messages(self):
        return list(ChatMessage.objects.filter(session__user=self.user).order_by("message_id"))
